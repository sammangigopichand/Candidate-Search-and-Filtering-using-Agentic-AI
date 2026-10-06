"""Recruiter dashboard for the Agentic AI recruitment system."""

import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent

if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))


import hashlib
import os
import tempfile
import smtplib
from email.message import EmailMessage
from dotenv import load_dotenv
from io import BytesIO

import streamlit as st
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select

from app.auth_store import (
    authenticate_user,
    auth_database_session,
    create_user,
    get_active_interviewers,
    update_interviewer_details,
)

from app.agents.jd_agent import process_job_description
from app.agents.coordinator_agent import (
    clear_workflow_checkpoints,
    process_candidate_workflow,
)
from app.agents.interview_agent import generate_interview_questions

from app.recruitment_store import (
    clear_candidate_recruitment_data,
    get_or_create_candidate,
    create_job,
    get_candidate_view_counts,
    get_candidate_view_for_job,
    get_candidate_by_resume_hash,
    list_active_jobs,
    recruitment_database_session,
    save_candidate_match,
)

from app.interview_store import (
    InterviewSession,
    clear_interview_processing_data,
    create_interview,
    database_session,
    find_existing_interview,
    get_interview,
    get_interview_details,
)

load_dotenv()


def send_interviewer_assignment_email(*, recipient, interviewer_name, candidate_name, job_title, interview_round):
    """Send an assignment email using the configured SMTP server."""
    host = os.getenv("SMTP_HOST", "").strip()
    username = os.getenv("SMTP_USERNAME", "").strip()
    password = os.getenv("SMTP_PASSWORD", "")
    sender = os.getenv("SMTP_FROM", "").strip() or username
    port = int(os.getenv("SMTP_PORT", "587"))
    missing = [
        key for key, value in (
            ("SMTP_HOST", host),
            ("SMTP_FROM or SMTP_USERNAME", sender),
        ) if not value
    ]
    if missing:
        raise RuntimeError(
            "Email is not configured. Set " + " and ".join(missing)
            + " in .env, then restart the recruiter app."
        )
    if username and not password:
        raise RuntimeError("SMTP_PASSWORD is required when SMTP_USERNAME is set.")
    if not recipient or "@" not in recipient:
        raise ValueError("The selected interviewer does not have a valid email address.")

    message = EmailMessage()
    message["Subject"] = f"Interview assignment: {candidate_name} · {job_title}"
    message["From"] = sender
    message["To"] = recipient
    message.set_content(
        f"Hello {interviewer_name},\n\n"
        f"You have been assigned to interview {candidate_name} for {job_title}.\n"
        f"Round: {interview_round}\n"
        "Sign in to the interviewer dashboard to review the assignment.\n"
    )

    use_ssl = os.getenv("SMTP_USE_SSL", "").strip().lower() in {"1", "true", "yes"} or port == 465
    smtp_class = smtplib.SMTP_SSL if use_ssl else smtplib.SMTP
    print(f"[email] Sending interview assignment to {recipient} via {host}:{port}.", flush=True)
    with smtp_class(host, port, timeout=20) as smtp:
        if not use_ssl and os.getenv("SMTP_STARTTLS", "true").strip().lower() in {"1", "true", "yes"}:
            smtp.starttls()
        if username:
            smtp.login(username, password)
        smtp.send_message(message)
    print(f"[email] Sent interview assignment to {recipient}.", flush=True)


# ============================================================
# PAGE CONFIGURATION
# ============================================================

def recruiter_login():
    st.title("AI Recruitment System")
    st.subheader("Recruiter login")
    with st.form("recruiter_login_form"):
        login = st.text_input("Username or email")
        password = st.text_input("Password", type="password")
        submitted = st.form_submit_button("Login", type="primary", width="stretch")
    if not submitted:
        return
    if not login.strip() or not password:
        st.error("Enter username/email and password.")
        return
    auth_factory = auth_database_session()
    auth_db = auth_factory()
    try:
        user = authenticate_user(auth_db, login, password)
        if user is None:
            st.error("Invalid username/email or password.")
        elif user.role != "recruiter":
            st.error("This account does not have recruiter access.")
        else:
            st.session_state.update({
                "authenticated": True,
                "auth_user_id": user.id,
                "auth_username": user.username,
                "auth_full_name": user.full_name,
                "auth_role": user.role,
            })
            st.rerun()
    finally:
        auth_db.close()


def recruiter_logout():
    for key in ("authenticated", "auth_user_id", "auth_username", "auth_full_name", "auth_role"):
        st.session_state.pop(key, None)
    st.rerun()


if not st.session_state.get("authenticated", False):
    recruiter_login()
    st.stop()
if st.session_state.get("auth_role") != "recruiter":
    st.error("You are not authorized to access the recruiter dashboard.")
    if st.button("Logout"):
        recruiter_logout()
    st.stop()

@st.cache_resource
def get_database_factory():
    return recruitment_database_session()


try:
    SessionLocal = get_database_factory()
except Exception as error:
    st.error(f"Database unavailable: {error}")
    st.stop()


sidebar_jobs = []
sidebar_pending_count = 0
sidebar_interview_rows = []
try:
    sidebar_db = SessionLocal()
    try:
        sidebar_jobs = list_active_jobs(sidebar_db)
        for job in sidebar_jobs:
            job_counts = get_candidate_view_counts(sidebar_db, job_id=job.id)
            sidebar_pending_count += job_counts.get("all", 0)
    finally:
        sidebar_db.close()
except Exception as error:
    st.warning(f"Could not load sidebar overview: {error}")

try:
    InterviewSessionFactory = database_session()
    sidebar_interview_db = InterviewSessionFactory()
    try:
        sidebar_interview_rows = sidebar_interview_db.execute(
            select(
                InterviewSession.job_id,
                InterviewSession.candidate_id,
                InterviewSession.status,
            )
        ).all()
    finally:
        sidebar_interview_db.close()
except Exception:
    # Interview filters remain available as "not scheduled" when there are
    # no interview records yet or interview storage is temporarily unavailable.
    sidebar_interview_rows = []

with st.sidebar:
    st.markdown("### AI Recruitment")
    st.caption(f"{st.session_state.get('auth_full_name', 'Recruiter')} · Recruiter")
    if st.session_state.get("auth_username"):
        st.caption(f"@{st.session_state['auth_username']}")
    st.divider()
    st.markdown("#### Quick overview")
    overview_col1, overview_col2 = st.columns(2)
    overview_col1.metric("Active jobs", len(sidebar_jobs))
    overview_col2.metric("Processed matches", sidebar_pending_count)

    sidebar_job_options = {
        f"{job.job_title} ({job.id[:8]})": job for job in sidebar_jobs
    }
    if sidebar_job_options:
        sidebar_selected_job = st.selectbox(
            "Filter by job",
            options=list(sidebar_job_options),
            key="candidate_review_job_sidebar",
        )
        sidebar_search = st.text_input(
            "Search candidates",
            placeholder="Name contains…",
            key="review_sidebar_search",
        )
        sidebar_status_filter = st.selectbox(
            "Interview status",
            ["All candidates", "Not scheduled", "In progress", "Completed"],
            key="review_sidebar_status",
        )
        sidebar_candidate_set = st.selectbox(
            "Candidate set",
            ["Relevant candidates", "All candidates"],
            key="review_sidebar_candidate_set",
        )
    else:
        sidebar_selected_job = None
        sidebar_search = ""
        sidebar_status_filter = "All candidates"
        sidebar_candidate_set = "Relevant candidates"
        st.caption("Save a job to enable candidate review filters.")

    st.divider()
    if st.button("Logout", width="stretch", key="recruiter_logout"):
        recruiter_logout()


for state_key, initial_value in (
    ("extracted_job_profile", None),
    ("raw_job_description", ""),
    ("saved_job_id", None),
):
    if state_key not in st.session_state:
        st.session_state[state_key] = initial_value


def display_list(title, values):
    st.markdown(f"**{title}**")
    if not values:
        st.caption("Not specified")
        return
    for value in values:
        st.write(f"- {value}")


def display_job_profile(profile):
    st.subheader("Extracted job requirements")
    left, right = st.columns(2)
    with left.container(border=True):
        st.markdown("#### Role information")
        st.write("**Job title:**", profile.get("job_title") or "Not specified")
        st.write("**Employment type:**", profile.get("employment_type") or "Not specified")
        st.write("**Location:**", profile.get("job_location") or "Not specified")
        st.write("**Experience:**", f"{profile.get('minimum_experience_years', 'Not specified')} – {profile.get('maximum_experience_years', 'Not specified')} years")
    with right.container(border=True):
        st.markdown("#### Role summary")
        st.write(profile.get("role_summary") or "No role summary was extracted.")
    left, right = st.columns(2)
    with left:
        display_list("Required skills", profile.get("required_skills") or [])
        display_list("Required certifications", profile.get("required_certifications") or [])
        display_list("Technologies", profile.get("technologies") or [])
    with right:
        display_list("Preferred skills", profile.get("preferred_skills") or [])
        display_list("Preferred certifications", profile.get("preferred_certifications") or [])
        display_list("Education", profile.get("education") or [])
        display_list("Responsibilities", profile.get("responsibilities") or [])
    with st.expander("Structured job profile"):
        st.json(profile)


def format_score(value):
    if value is None or (isinstance(value, str) and value.lower().strip() == "not assessed"):
        return "Not Assessed"
    try:
        return f"{float(value):.2f}%"
    except (TypeError, ValueError):
        return str(value)


def classify_experience_years(years):
    try:
        years = float(years)
    except (TypeError, ValueError):
        return None
    if years <= 0:
        return "Fresher"
    if years < 2:
        return "Junior"
    if years < 5:
        return "Mid-Level"
    return "Senior"


def get_experience_display(report):
    report = report or {}
    total = report.get("total_experience_years") or {}
    low = classify_experience_years(total.get("minimum"))
    high = classify_experience_years(total.get("maximum"))
    provisional = report.get("provisional_experience_level") or "Undetermined"
    if low and high and low != high:
        return f"{low} – {high}", "Date-dependent"
    if provisional != "Undetermined":
        return provisional, None
    return low or "Undetermined", None


def format_experience_range(value):
    value = value or {}
    minimum, maximum = value.get("minimum"), value.get("maximum")
    if minimum is None or maximum is None:
        return "Not verified"
    return f"{minimum} years" if minimum == maximum else f"{minimum} - {maximum} years"


def render_employment_evidence(report):
    roles = report.get("assessed_roles") or []
    unverified = report.get("unverified_roles") or []
    st.subheader("Employment evidence")
    if not roles and not unverified:
        st.caption("No professional employment evidence was available.")
        return
    for role in roles:
        if not isinstance(role, dict):
            st.write(role)
            continue
        title = role.get("job_title") or role.get("role") or "Role not specified"
        company = role.get("company") or "Company not specified"
        st.markdown(f"**{title} — {company}**")
        st.caption(f"{role.get('start_date') or 'Unknown'} → {role.get('end_date') or 'Unknown'}")
        relevance = role.get("relevance") or {}
        for evidence in relevance.get("evidence", []) if isinstance(relevance, dict) else []:
            st.write(f"• {evidence}")
    if unverified:
        with st.expander("Unverified employment records"):
            st.write(unverified)


def get_existing_interview_info(candidate_id, job_id, interview_round):
    factory = database_session()
    db = factory()
    try:
        record = find_existing_interview(db, candidate_id=candidate_id, job_id=job_id, interview_round=interview_round)
        if not record:
            return None
        return {"id": record.id, "interviewer_id": record.interviewer_id, "status": record.status,
                "interview_round": record.interview_round, "interview_date": record.interview_date}
    finally:
        db.close()


def get_evaluation_report(session_id, interviewer_id):
    factory = database_session()
    db = factory()
    try:
        record = get_interview(db, session_id, interviewer_id)
        return dict(record.report.data or {}) if record and record.report else None
    finally:
        db.close()


def render_interview_evaluation(report):
    if not report:
        st.warning("The interview is marked completed, but no evaluation report was found.")
        return
    st.markdown("### Interview evaluation")
    cols = st.columns(4)
    for col, label, key in zip(cols, ("Overall", "Technical", "Project", "Behavioral"),
                               ("overall_score", "technical_score", "project_score", "behavioral_score")):
        col.metric(label, format_score(report.get(key)))
    with st.expander("View interview evaluation"):
        for key, label in (("observed_strengths", "Observed strengths"), ("partially_met", "Partially demonstrated"),
                           ("areas_requiring_improvement", "Areas to improve"), ("unassessed_criteria", "Not assessed")):
            st.markdown(f"#### {label}")
            values = report.get(key) or []
            if values:
                for value in values:
                    st.write(f"• {value.get('criterion') or value.get('text') or value if isinstance(value, dict) else value}")
            else:
                st.caption("No criteria recorded.")
        st.caption(report.get("score_basis", "Scores are based on interviewer checklist assessments."))


with st.container(horizontal=True, gap="small"):
    for number, label, detail in (
        ("01", "Define role", "Create and save a job"),
        ("02", "Build shortlist", "Process candidate resumes"),
        ("03", "Review", "Compare evidence and plan interviews"),
    ):
        with st.container(border=True):
            st.caption(f"STEP {number}")
            st.markdown(f"**{label}**")
            st.caption(detail)


# ============================================================
# TABS
# ============================================================

def make_interview_performance_workbook(summary_rows, detail_rows):
    workbook = Workbook()
    header_fill = PatternFill("solid", fgColor="176B68")
    header_font = Font(color="FFFFFF", bold=True)

    def add_sheet(title, rows):
        sheet = workbook.create_sheet(title=title)
        if not rows:
            sheet.append(["No completed interview data available."])
            return
        headers = list(rows[0])
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(header, "") for header in headers])
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 48)
            sheet.column_dimensions[column[0].column_letter].width = width

    add_sheet("All Job Profiles", summary_rows)
    add_sheet("Checklist Details", detail_rows)

    groups = {}
    for row in summary_rows:
        group_key = (row["Job ID"], row["Job Profile"])
        groups.setdefault(group_key, []).append(row)

    used_names = {"All Job Profiles", "Checklist Details"}
    for (job_id, job_title), rows in groups.items():
        base_name = "".join(char for char in str(job_title) if char not in "[]:*?/\\")
        base_name = (base_name.strip() or "Job Profile")[:25]
        sheet_name = base_name
        suffix = 1
        while sheet_name in used_names:
            suffix += 1
            sheet_name = f"{base_name[:20]} ({suffix})"
        used_names.add(sheet_name)
        add_sheet(sheet_name, rows)

    # Replace the empty default worksheet created by openpyxl.
    default_sheet = workbook["Sheet"]
    workbook.remove(default_sheet)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()

# create_tab, jobs_tab, candidates_tab, review_tab = st.tabs(
#     [
#         "Create Job",
#         "Active Jobs",
#         "Process Candidates",
#         "Candidate Review",
#     ]
# )


(
    create_tab,
    jobs_tab,
    candidates_tab,
    review_tab,
    reports_tab,
    user_management_tab,
) = st.tabs(
    [
        ":material/add_business: Create job",
        ":material/work: Active jobs",
        ":material/person_search: Process candidates",
        ":material/rate_review: Candidate review",
        ":material/monitoring: Interview reports",
        ":material/manage_accounts: User management",
    ]
)


# ============================================================
# CREATE JOB
# ============================================================

with create_tab:

    st.header("Create a job", icon=":material/post_add:")

    st.write(
        "Enter the complete job description. "
        "The Job Description Understanding Agent will extract "
        "structured hiring requirements."
    )

    job_description = st.text_area(
        "Job Description",
        value=st.session_state.get("raw_job_description", ""),
        height=300,
        placeholder=(
            "Example:\n\n"
            "We are hiring a Data Scientist with experience "
            "in Python, machine learning, Pandas and "
            "Scikit-learn..."
        ),
    )

    process_button = st.button(
        "Process Job Description",
        type="primary",
        use_container_width=True,
    )

    if process_button:

        if not job_description.strip():

            st.warning(
                "Enter a job description before processing."
            )

        else:

            try:

                with st.spinner(
                    "Job Description Agent is processing "
                    "the requirements..."
                ):

                    profile = process_job_description(
                        job_description
                    )

                st.session_state.extracted_job_profile = profile
                st.session_state.raw_job_description = (
                    job_description
                )

                st.session_state.saved_job_id = None

                st.success(
                    "Job description processed successfully."
                )

            except Exception as error:

                st.error(
                    f"Could not process job description: {error}"
                )

    profile = st.session_state.extracted_job_profile

    if profile:

        st.divider()

        display_job_profile(profile)

        st.divider()

        st.subheader("Save Job")

        st.info(
            "Review the extracted requirements before saving. "
            "The recruiter remains responsible for verifying "
            "the job requirements."
        )

        if st.session_state.saved_job_id:

            st.success(
                "This job has already been saved.\n\n"
                f"Job ID: {st.session_state.saved_job_id}"
            )

        else:

            save_button = st.button(
                "Save Job to Database",
                use_container_width=True,
            )

            if save_button:

                try:

                    db = SessionLocal()

                    try:

                        job_id = create_job(
                            db,
                            raw_job_description=(
                                st.session_state
                                .raw_job_description
                            ),
                            job_profile=profile,
                        )

                    finally:
                        db.close()

                    st.session_state.saved_job_id = job_id

                    st.success(
                        "Job saved successfully."
                    )

                    st.write(
                        "**Job ID:**",
                        job_id,
                    )

                except Exception as error:

                    st.error(
                        f"Could not save job: {error}"
                    )


# ============================================================
# ACTIVE JOBS
# ============================================================

with jobs_tab:

    st.header("Active Jobs")

    try:

        db = SessionLocal()

        try:
            jobs = list_active_jobs(db)
        finally:
            db.close()

        if not jobs:

            st.info(
                "No active jobs are available yet."
            )

        else:

            st.write(
                f"Active jobs: **{len(jobs)}**"
            )

            for job in jobs:

                created_date = (
                    job.created_at.strftime("%Y-%m-%d")
                    if job.created_at
                    else "Unknown date"
                )

                title = (
                    f"{job.job_title} | {created_date}"
                )

                with st.expander(title):

                    st.write(
                        "**Job ID:**",
                        job.id,
                    )

                    st.write(
                        "**Status:**",
                        job.status,
                    )

                    profile = job.job_profile or {}

                    required = (
                        profile.get("required_skills")
                        or []
                    )

                    preferred = (
                        profile.get("preferred_skills")
                        or []
                    )

                    st.write(
                        "**Required Skills:**",
                        (
                            ", ".join(required)
                            if required
                            else "Not specified"
                        ),
                    )

                    st.write(
                        "**Preferred Skills:**",
                        (
                            ", ".join(preferred)
                            if preferred
                            else "Not specified"
                        ),
                    )

                    minimum = profile.get(
                        "minimum_experience_years"
                    )

                    st.write(
                        "**Minimum Experience:**",
                        (
                            f"{minimum} years"
                            if minimum is not None
                            else "Not specified"
                        ),
                    )

                    with st.expander(
                        "View Structured Job Profile"
                    ):
                        st.json(profile)

    except Exception as error:

        st.error(
            f"Could not load active jobs: {error}"
        )


# ============================================================
# PROCESS CANDIDATES
# ============================================================

with candidates_tab:

    st.header("Build your candidate shortlist", icon=":material/person_search:")

    # ============================================================
    # CANDIDATE DATA MANAGEMENT
    # ============================================================

    with st.expander(
        "Candidate Data Management",
        expanded=False,
    ):

        st.warning(
            "This permanently deletes all candidate processing "
            "data, including candidates, candidate-job matches, "
            "interview sessions, evaluation reports and LangGraph "
            "workflow history. Saved jobs will be preserved."
        )

        confirm_clear_data = st.checkbox(
            "I understand that this action cannot be undone.",
            key="confirm_clear_candidate_processing_data",
        )

        clear_candidate_data_button = st.button(
            "Clear All Candidate Processing Data",
            disabled=not confirm_clear_data,
            use_container_width=True,
            key="clear_all_candidate_processing_data",
        )

        if clear_candidate_data_button:

            try:

                # ----------------------------------------------
                # 1. DELETE INTERVIEW / EVALUATION DATA
                # ----------------------------------------------

                InterviewSessionFactory = database_session()
                interview_db = InterviewSessionFactory()

                try:

                    clear_interview_processing_data(
                        interview_db
                    )

                finally:

                    interview_db.close()

                # ----------------------------------------------
                # 2. DELETE CANDIDATE / MATCH DATA
                # ----------------------------------------------

                recruitment_db = SessionLocal()

                try:

                    clear_candidate_recruitment_data(
                        recruitment_db
                    )

                finally:

                    recruitment_db.close()

                # ----------------------------------------------
                # 3. DELETE LANGGRAPH WORKFLOW HISTORY
                # ----------------------------------------------

                clear_workflow_checkpoints()

                st.success(
                    "All candidate processing data was deleted. "
                    "Saved jobs were preserved."
                )

                st.rerun()


            except Exception as error:

                st.error(
                    "Could not clear candidate processing data: "
                    f"{error}"
                )

    st.divider()

    st.write(
        "Choose a saved role, add resumes, then review each candidate's "
        "skills and experience evidence. Results are linked to the selected job."
    )


    try:

        db = SessionLocal()

        try:
            available_jobs = list_active_jobs(db)
        finally:
            db.close()

    except Exception as error:

        st.error(
            f"Could not load jobs: {error}"
        )

        available_jobs = []

    if not available_jobs:

        st.info(
            "Create and save a job before processing candidates."
        )

    else:

        job_lookup = {
            (
                f"{job.job_title} | "
                f"{job.id[:8]}"
            ): job
            for job in available_jobs
        }

        selected_job_label = st.selectbox(
            "Select Job",
            options=list(job_lookup.keys()),
            key="process_candidate_job",
        )

        selected_job = job_lookup[
            selected_job_label
        ]

        job_profile = selected_job.job_profile or {}

        st.markdown("### Role selected")

        col1, col2, col3 = st.columns(3)

        col1.metric(
            "Job",
            selected_job.job_title,
        )

        required_skills = (
            job_profile.get("required_skills")
            or []
        )

        col2.metric(
            "Required Skills",
            len(required_skills),
        )

        minimum_experience = (
            job_profile.get(
                "minimum_experience_years"
            )
        )

        col3.metric(
            "Minimum Experience",
            (
                f"{minimum_experience} years"
                if minimum_experience is not None
                else "Not specified"
            ),
        )

        with st.expander(
            "View Job Requirements"
        ):

            display_job_profile(
                job_profile
            )

        st.divider()

        st.subheader(
            "Upload Candidate Resumes"
        )

        uploaded_resumes = st.file_uploader(
            "Upload PDF or DOCX resumes",
            type=[
                "pdf",
                "docx",
            ],
            accept_multiple_files=True,
            key="candidate_resume_upload",
        )

        if uploaded_resumes:

            st.caption(f"{len(uploaded_resumes)} resume(s) ready to process")
            with st.expander("Review uploaded files", expanded=False):
                for uploaded_file in uploaded_resumes:
                    st.write(f":material/description: {uploaded_file.name}")

        process_candidates_button = st.button(
            "Process Candidates",
            type="primary",
            use_container_width=True,
        )

        if process_candidates_button:

            if not uploaded_resumes:

                st.warning(
                    "Upload at least one resume."
                )

            else:

                st.divider()

                st.subheader(
                    "Processing results",
                    icon=":material/analytics:",
                )

                successful = 0
                failed = 0

                progress_bar = st.progress(0)
                status_text = st.empty()

                total_files = len(
                    uploaded_resumes
                )

                for index, uploaded_file in enumerate(
                    uploaded_resumes,
                    start=1,
                ):

                    status_text.write(
                        f"Processing "
                        f"{uploaded_file.name} "
                        f"({index}/{total_files})..."
                    )

                    temporary_path = None

                    try:

                        # ==========================================
                        # RESUME HASH
                        # ==========================================

                        resume_bytes = uploaded_file.getvalue()

                        resume_hash = hashlib.sha256(
                            resume_bytes
                        ).hexdigest()

                        extension = os.path.splitext(
                            uploaded_file.name
                        )[1].lower()

                        with tempfile.NamedTemporaryFile(
                            delete=False,
                            suffix=extension,
                        ) as temporary_file:

                            temporary_file.write(
                                resume_bytes
                            )

                            temporary_path = (
                                temporary_file.name
                            )

                        existing_candidate_db = SessionLocal()
                        try:
                            existing_candidate = get_candidate_by_resume_hash(
                                existing_candidate_db,
                                resume_hash,
                            )
                            cached_candidate_profile = (
                                dict(existing_candidate.candidate_profile or {})
                                if existing_candidate
                                else None
                            )
                        finally:
                            existing_candidate_db.close()

                        # ==========================================
                        # LANGGRAPH COORDINATOR
                        # ==========================================

                        workflow_result = (
                            process_candidate_workflow(
                                resume_path=temporary_path,
                                job_profile=job_profile,
                                candidate_profile=cached_candidate_profile,
                            )
                        )



                        # ==========================================
                        # LANGGRAPH WORKFLOW ID
                        # ==========================================

                        workflow_id = workflow_result.get(
                            "workflow_id"
                        )

                        if not workflow_id:

                            raise RuntimeError(
                                "Coordinator did not return a workflow_id."
                            )

                        # ==========================================
                        # VALIDATE WORKFLOW
                        # ==========================================

                        workflow_error = workflow_result.get(
                            "error"
                        )

                        if workflow_error:

                            raise RuntimeError(
                                "Coordinator workflow failed: "
                                f"{workflow_error}"
                            )

                        workflow_status = workflow_result.get(
                            "workflow_status"
                        )

                        current_stage = workflow_result.get(
                            "current_stage"
                        )

                        if (
                            workflow_status
                            != "awaiting_recruiter_review"
                            or current_stage
                            != "recruiter_review"
                        ):

                            raise RuntimeError(
                                "Coordinator workflow did not reach "
                                "recruiter review. "
                                f"Stage: {current_stage}, "
                                f"Status: {workflow_status}"
                            )

                        # ==========================================
                        # COORDINATOR OUTPUTS
                        # ==========================================

                        candidate_profile = workflow_result.get(
                            "candidate_profile"
                        )

                        matching_report = workflow_result.get(
                            "matching_report"
                        )

                        experience_report = workflow_result.get(
                            "experience_report"
                        )

                        if not candidate_profile:

                            raise ValueError(
                                "Coordinator returned no candidate profile."
                            )

                        if not matching_report:

                            raise ValueError(
                                "Coordinator returned no matching report."
                            )

                        if not experience_report:

                            raise ValueError(
                                "Coordinator returned no experience report."
                            )

                        # ==========================================
                        # SAVE OR REUSE CANDIDATE
                        # ==========================================

                        db = SessionLocal()

                        try:

                            candidate_result = (
                                get_or_create_candidate(
                                    db,
                                    candidate_profile=(
                                        candidate_profile
                                    ),
                                    resume_filename=(
                                        uploaded_file.name
                                    ),
                                    resume_hash=resume_hash,
                                )
                            )

                            candidate_id = (
                                candidate_result[
                                    "candidate_id"
                                ]
                            )

                            candidate_created = (
                                candidate_result[
                                    "created"
                                ]
                            )

                            candidate_reused = (
                                candidate_result[
                                    "reused"
                                ]
                            )

                            # ======================================
                            # SAVE / UPDATE JOB MATCH
                            # ======================================

                            match_id = save_candidate_match(
                                db,
                                candidate_id=candidate_id,
                                job_id=selected_job.id,
                                matching_report=(
                                    matching_report
                                ),
                                experience_report=(
                                    experience_report
                                ),
                                workflow_id=workflow_id,
                            )

                            st.write(
                                "**Workflow ID:**",
                                workflow_id,
                            )
                        finally:
                            db.close()

                        successful += 1

                        candidate_name = (
                            candidate_profile.get("name")
                            or "Unknown Candidate"
                        )

                        if candidate_reused:

                            st.info(
                                f"{candidate_name} already exists "
                                "in the candidate database. "
                                "The existing candidate record "
                                "was reused and the job match "
                                "was updated."
                            )

                        elif candidate_created:

                            st.success(
                                f"{candidate_name} was added "
                                "as a new candidate."
                            )

                        required_coverage = (
                            matching_report
                            .get(
                                "required_skills",
                                {},
                            )
                            .get(
                                "coverage_percentage"
                            )
                        )

                        (
                            experience_level,
                            experience_note,
                        ) = get_experience_display(
                            experience_report
                        )

                        with st.expander(
                            f"{candidate_name} - "
                            f"Processed Successfully",
                            expanded=True,
                        ):

                            col1, col2, col3 = (
                                st.columns(3)
                            )

                            col1.metric(
                                "Candidate",
                                candidate_name,
                            )

                            col2.metric(
                                "Required Skill Coverage",
                                (
                                    f"{required_coverage}%"
                                    if required_coverage
                                    is not None
                                    else "Not Assessed"
                                ),
                            )

                            col3.metric(
                                "Experience Level",
                                experience_level,
                            )

                            if experience_note:
                                col3.caption(
                                    experience_note
                                )

                            st.write(
                                "**Candidate ID:**",
                                candidate_id,
                            )

                            st.write(
                                "**Match ID:**",
                                match_id,
                            )

                            matched_skills = (
                                matching_report
                                .get(
                                    "required_skills",
                                    {},
                                )
                                .get(
                                    "matched",
                                    [],
                                )
                            )

                            missing_skills = (
                                matching_report
                                .get(
                                    "required_skills",
                                    {},
                                )
                                .get(
                                    "missing",
                                    [],
                                )
                            )

                            col1, col2 = (
                                st.columns(2)
                            )

                            with col1:

                                st.markdown(
                                    "#### Matched Required Skills"
                                )

                                if matched_skills:

                                    for item in matched_skills:

                                        if isinstance(
                                            item,
                                            dict,
                                        ):

                                            skill = (
                                                item.get(
                                                    "required_skill"
                                                )
                                                or str(item)
                                            )

                                        else:
                                            skill = str(item)

                                        st.write(
                                            f"✓ {skill}"
                                        )

                                else:

                                    st.caption(
                                        "No required skill "
                                        "matches were documented."
                                    )

                            with col2:

                                st.markdown(
                                    "#### Missing Required Skills"
                                )

                                if missing_skills:

                                    for skill in missing_skills:

                                        st.write(
                                            f"• {skill}"
                                        )

                                else:

                                    st.caption(
                                        "No required skills "
                                        "were identified as missing."
                                    )

                            st.markdown(
                                "#### Experience Assessment"
                            )

                            total_experience = (
                                experience_report.get(
                                    "total_experience_years",
                                    {},
                                )
                            )

                            relevant_experience = (
                                experience_report.get(
                                    "documented_relevant_"
                                    "experience_years",
                                    {},
                                )
                            )

                            col1, col2 = (
                                st.columns(2)
                            )

                            col1.write(
                                "**Documented Total Experience:**"
                            )

                            col1.write(
                                f"{total_experience.get('minimum', 0)} "
                                f"- "
                                f"{total_experience.get('maximum', 0)} "
                                f"years"
                            )

                            col2.write(
                                "**Documented Relevant Experience:**"
                            )

                            col2.write(
                                f"{relevant_experience.get('minimum', 0)} "
                                f"- "
                                f"{relevant_experience.get('maximum', 0)} "
                                f"years"
                            )

                            with st.expander(
                                "Candidate Profile"
                            ):
                                st.json(
                                    candidate_profile
                                )

                            with st.expander(
                                "Matching Report"
                            ):
                                st.json(
                                    matching_report
                                )

                            with st.expander(
                                "Experience Report"
                            ):
                                st.json(
                                    experience_report
                                )

                    except Exception as error:

                        failed += 1

                        st.error(
                            f"{uploaded_file.name}: "
                            f"{error}"
                        )

                    finally:

                        if (
                            temporary_path
                            and os.path.exists(
                                temporary_path
                            )
                        ):

                            try:
                                os.remove(
                                    temporary_path
                                )
                            except OSError:
                                pass

                    progress_bar.progress(
                        index / total_files
                    )

                status_text.success(
                    "Candidate processing completed."
                )

                st.divider()

                col1, col2, col3 = st.columns(3)

                col1.metric(
                    "Uploaded",
                    total_files,
                )

                col2.metric(
                    "Processed Successfully",
                    successful,
                )

                col3.metric(
                    "Failed",
                    failed,
                )


# ============================================================
# CANDIDATE REVIEW
# ============================================================

with review_tab:

    st.header("Candidate Review")

    email_notice = st.session_state.pop("assignment_email_notice", None)
    if email_notice:
        notice_type, notice_message = email_notice
        getattr(st, notice_type)(notice_message)

    st.write(
        "Review candidate evidence before deciding whether "
        "the candidate should proceed to interview."
    )

    # ========================================================
    # INTERVIEW SETTINGS
    # ========================================================

    st.caption("Choose an interviewer from the candidate card when you are ready to assign.")

    try:
        AuthSessionLocal = auth_database_session()
        auth_db = AuthSessionLocal()
        try:
            active_interviewers = get_active_interviewers(auth_db)
        finally:
            auth_db.close()
    except Exception as error:
        active_interviewers = []
        st.error(f"Could not load interviewers: {error}")

    if not active_interviewers:
        st.warning("No active interviewers are available. Create an interviewer from User Management.")

    review_db = SessionLocal()

    try:

        jobs = list_active_jobs(
            review_db
        )

        if not jobs:

            st.info(
                "No active jobs are available."
            )

        else:

            job_options = {
                (
                    f"{job.job_title} "
                    f"({job.id[:8]})"
                ): job
                for job in jobs
            }

            selected_job_label = st.session_state.get(
                "candidate_review_job_sidebar"
            )
            if selected_job_label not in job_options:
                selected_job_label = next(iter(job_options))

            selected_job = job_options[
                selected_job_label
            ]

            st.divider()

            st.subheader(
                selected_job.job_title
            )

            # =================================================
            # CANDIDATE VIEW
            # =================================================

            st.subheader("Candidate view")

            view_mode = (
                "All Candidates"
                if sidebar_candidate_set == "All candidates"
                else "Relevant Candidates"
            )
            st.caption(
                f"Sidebar filters · {view_mode} · "
                f"{sidebar_status_filter} · "
                f"{sidebar_search.strip() or 'All names'}"
            )

            minimum_skill_coverage = 60.0
            require_relevant_experience = False
            maximum_missing_required_skills = None

            if view_mode == "Relevant Candidates":

                st.markdown(
                    "#### Relevance Filters"
                )

                filter_col1, filter_col2 = (
                    st.columns(2)
                )

                with filter_col1:

                    minimum_skill_coverage = st.slider(
                        "Minimum Required-Skill Coverage",
                        min_value=0,
                        max_value=100,
                        value=60,
                        step=5,
                        format="%d%%",
                        key="review_minimum_skill_coverage",
                        help=(
                            "Candidates below this documented "
                            "required-skill coverage are hidden "
                            "from the Relevant Candidates view "
                            "only. They remain stored and are "
                            "available under All Candidates."
                        ),
                    )

                    require_relevant_experience = (
                        st.checkbox(
                            "Require documented relevant "
                            "professional experience",
                            value=False,
                            key=(
                                "review_require_"
                                "relevant_experience"
                            ),
                            help=(
                                "When enabled, the Experience "
                                "Agent must have established "
                                "some job-relevant professional "
                                "employment evidence."
                            ),
                        )
                    )

                with filter_col2:

                    limit_missing_skills = st.checkbox(
                        "Limit missing required skills",
                        value=False,
                        key="review_limit_missing_skills",
                    )

                    if limit_missing_skills:

                        maximum_missing_required_skills = (
                            st.number_input(
                                "Maximum Missing Required Skills",
                                min_value=0,
                                value=2,
                                step=1,
                                key=(
                                    "review_maximum_"
                                    "missing_skills"
                                ),
                            )
                        )

                st.caption(
                    "These filters control the default recruiter "
                    "view only. They do not delete, reject, rank, "
                    "or make a hiring decision about any candidate."
                )

            normalized_view = (
                "all"
                if view_mode == "All Candidates"
                else "relevant"
            )

            matches = get_candidate_view_for_job(
                review_db,
                job_id=selected_job.id,
                view=normalized_view,
                minimum_skill_coverage=(
                    minimum_skill_coverage
                ),
                require_relevant_experience=(
                    require_relevant_experience
                ),
                maximum_missing_required_skills=(
                    maximum_missing_required_skills
                ),
            )

            counts = get_candidate_view_counts(
                review_db,
                job_id=selected_job.id,
                minimum_skill_coverage=(
                    minimum_skill_coverage
                ),
                require_relevant_experience=(
                    require_relevant_experience
                ),
                maximum_missing_required_skills=(
                    maximum_missing_required_skills
                ),
            )

            total_before_sidebar_filters = len(matches)
            interview_status_by_candidate = {}
            for job_id, candidate_id, interview_status in sidebar_interview_rows:
                if job_id != selected_job.id:
                    continue
                current_status = interview_status_by_candidate.get(candidate_id)
                if interview_status == "submitted" or current_status is None:
                    interview_status_by_candidate[candidate_id] = interview_status

            search_term = sidebar_search.strip().casefold()
            filtered_matches = []
            for match_record, candidate in matches:
                profile = candidate.candidate_profile or {}
                candidate_name = (
                    candidate.name
                    or profile.get("name")
                    or "Unknown Candidate"
                )
                if search_term and search_term not in candidate_name.casefold():
                    continue

                existing_interview_status = interview_status_by_candidate.get(
                    candidate.id
                )
                if sidebar_status_filter == "Not scheduled" and existing_interview_status:
                    continue
                if sidebar_status_filter == "In progress" and (
                    not existing_interview_status
                    or existing_interview_status == "submitted"
                ):
                    continue
                if sidebar_status_filter == "Completed" and existing_interview_status != "submitted":
                    continue
                filtered_matches.append((match_record, candidate))

            matches = filtered_matches
            def candidate_match_score(row):
                report = row[0].matching_report or {}
                required = report.get("required_skills") or {}
                try:
                    return float(required.get("coverage_percentage") or 0)
                except (TypeError, ValueError):
                    return 0.0

            def candidate_match_level(score):
                if score >= 80:
                    return "High Match"
                if score >= 60:
                    return "Moderate Match"
                return "Low Match"

            matches.sort(key=candidate_match_score, reverse=True)
            st.caption(
                f"{len(matches)} candidate(s) match these sidebar filters "
                f"out of {total_before_sidebar_filters} in this candidate set."
            )

            if normalized_view == "relevant":

                st.caption(
                    f"Showing {counts['relevant']} of "
                    f"{counts['all']} processed candidate(s). "
                    f"{counts['hidden_by_current_filter']} "
                    "candidate(s) are hidden by the current "
                    "relevance filters."
                )

            else:

                st.caption(
                    f"Showing all {counts['all']} "
                    "processed candidate(s)."
                )

            if not matches:

                if counts["all"] == 0:

                    st.info(
                        "No candidates have been processed "
                        "for this job yet."
                    )

                else:

                    st.info(
                        "No candidates satisfy the current "
                        "relevance filters. Adjust the filters "
                        "or select All Candidates to review "
                        "every processed candidate."
                    )

            else:

                # =============================================
                # INTERACTIVE MATCH TABLE
                # =============================================

                table_rows = []
                for candidate_index, (match_record, candidate) in enumerate(matches, start=1):
                    profile = candidate.candidate_profile or {}
                    score = candidate_match_score((match_record, candidate))
                    experience_level, _ = get_experience_display(
                        match_record.experience_report or {}
                    )
                    table_rows.append({
                        "No.": candidate_index,
                        "Candidate": candidate.name or profile.get("name") or "Unknown Candidate",
                        "Matching Level": candidate_match_level(score),
                        "Match Score": score,
                        "Experience": experience_level,
                        "Interview Status": (
                            "Completed"
                            if interview_status_by_candidate.get(candidate.id) == "submitted"
                            else "In progress"
                            if interview_status_by_candidate.get(candidate.id)
                            else "Not scheduled"
                        ),
                    })

                st.caption("Sorted by required-skill match score, highest first. Click View Details to open a candidate review.")
                profile_button_key = f"open_candidate_profile_{selected_job.id}"
                selected_index_key = f"selected_candidate_index_{selected_job.id}"

                def open_candidate_profile():
                    click_state = st.session_state.get(profile_button_key)
                    if click_state is not None and click_state.row is not None:
                        if st.session_state.get(selected_index_key) == click_state.row:
                            st.session_state.pop(selected_index_key, None)
                        else:
                            st.session_state[selected_index_key] = click_state.row

                for row in table_rows:
                    row["Details"] = "View Details"

                st.dataframe(
                    pd.DataFrame(table_rows),
                    column_config={
                        "Match Score": st.column_config.ProgressColumn(
                            "Match Score", min_value=0, max_value=100, format="%.0f%%"
                        ),
                        "Details": st.column_config.ButtonColumn(
                            "Details", on_click=open_candidate_profile, key=profile_button_key
                        ),
                    },
                    hide_index=True,
                    key=f"candidate_match_table_{selected_job.id}",
                )
                selected_index = st.session_state.get(selected_index_key)
                if selected_index is not None:
                    selected_index = min(selected_index, len(matches) - 1)
                    st.session_state[selected_index_key] = selected_index
                    matches = [matches[selected_index]]
                    st.divider()
                    st.subheader("Selected Candidate Review")
                else:
                    st.info("Choose a candidate with View Details to open their profile and review.")
                    matches = []

                # =============================================
                # SELECTED CANDIDATE DETAIL
                # =============================================

                for match_record, candidate in matches:

                    profile = (
                        candidate.candidate_profile
                        or {}
                    )

                    matching_report = (
                        match_record.matching_report
                        or {}
                    )

                    experience_report = (
                        match_record.experience_report
                        or {}
                    )

                    candidate_name = (
                        candidate.name
                        or profile.get("name")
                        or "Unknown Candidate"
                    )

                    required_match = (
                        matching_report.get(
                            "required_skills",
                            {},
                        )
                    )

                    coverage = (
                        required_match.get(
                            "coverage_percentage"
                        )
                    )

                    (
                        experience_level,
                        experience_note,
                    ) = get_experience_display(
                        experience_report
                    )

                    candidate_id = candidate.id
                    job_id = selected_job.id
                    match_id = match_record.id
                    assignment_key = f"assignment_{candidate_id}_{job_id}"
                    assignment = st.session_state.get(assignment_key)
                    interviewer_id = assignment.get("interviewer_id", "") if assignment else ""
                    interview_round = assignment.get("interview_round", "Technical Round") if assignment else "Technical Round"

                    job_title = (
                        selected_job.job_title
                    )

                    job_profile = dict(
                        selected_job.job_profile
                        or {}
                    )

                    with st.container(border=True):

                        st.subheader(f"Candidate Profile · {candidate_name}")
                        profile_col1, profile_col2 = st.columns(2)
                        profile_col1.markdown("**Contact**")
                        profile_col1.write(f"Email: {profile.get('email') or 'Not provided'}")
                        profile_col1.write(f"Phone: {profile.get('phone') or profile.get('phone_number') or 'Not provided'}")
                        profile_col1.write(f"Location: {profile.get('location') or profile.get('address') or 'Not provided'}")
                        profile_col2.markdown("**Profile Summary**")
                        profile_col2.write(
                            profile.get("professional_summary")
                            or profile.get("summary")
                            or profile.get("career_objective")
                            or "No profile summary was extracted."
                        )

                        # =====================================
                        # SUMMARY
                        # =====================================

                        col1, col2, col3, col4 = st.columns(4)

                        match_score = candidate_match_score((match_record, candidate))
                        match_level = candidate_match_level(match_score)
                        col1.metric("Matching Level", match_level)
                        col2.metric("Match Score", f"{match_score:.0f}%")

                        col3.metric(
                            "Experience Level",
                            experience_level,
                        )

                        if experience_note:
                            col3.caption(
                                experience_note
                            )

                        col4.metric(
                            "Candidate",
                            candidate_name,
                        )

                        st.caption(
                            f"Candidate ID: "
                            f"{candidate_id}"
                        )

                        st.divider()

                        # =====================================
                        # MATCHED SKILLS
                        # =====================================

                        st.subheader(
                            "Matched Required Skills"
                        )

                        matched_skills = (
                            required_match.get(
                                "matched",
                                [],
                            )
                        )

                        if matched_skills:

                            for skill in matched_skills:

                                if isinstance(
                                    skill,
                                    dict,
                                ):

                                    required_skill = (
                                        skill.get(
                                            "required_skill"
                                        )
                                    )

                                    candidate_skill = (
                                        skill.get(
                                            "candidate_skill"
                                        )
                                    )

                                    if (
                                        candidate_skill
                                        and candidate_skill
                                        != required_skill
                                    ):

                                        st.write(
                                            f"✓ {required_skill} "
                                            f"— matched with "
                                            f"{candidate_skill}"
                                        )

                                    else:

                                        st.write(
                                            f"✓ {required_skill}"
                                        )

                                else:

                                    st.write(
                                        f"✓ {skill}"
                                    )

                        else:

                            st.write(
                                "No required skill matches "
                                "were documented."
                            )

                        # =====================================
                        # MISSING SKILLS
                        # =====================================

                        st.subheader(
                            "Missing Required Skills"
                        )

                        missing_skills = (
                            required_match.get(
                                "missing",
                                [],
                            )
                        )

                        if missing_skills:

                            for skill in missing_skills:

                                st.write(
                                    f"• {skill}"
                                )

                        else:

                            st.write(
                                "No required skills are "
                                "currently missing."
                            )

                        # =====================================
                        # EXPERIENCE
                        # =====================================

                        st.subheader(
                            "Experience Evidence"
                        )

                        total_experience = (
                            experience_report.get(
                                "total_experience_years",
                                {},
                            )
                        )

                        relevant_experience = (
                            experience_report.get(
                                "documented_relevant_"
                                "experience_years",
                                {},
                            )
                        )

                        total_text = (
                            format_experience_range(
                                total_experience
                            )
                        )

                        relevant_text = (
                            format_experience_range(
                                relevant_experience
                            )
                        )

                        exp_col1, exp_col2 = (
                            st.columns(2)
                        )

                        exp_col1.metric(
                            "Total Experience",
                            total_text,
                        )

                        exp_col2.metric(
                            "Relevant Experience",
                            relevant_text,
                        )

                        render_employment_evidence(
                            experience_report
                        )

                        # =====================================
                        # EDUCATION
                        # =====================================

                        st.subheader("Education")

                        education = (
                            profile.get(
                                "education",
                                [],
                            )
                        )

                        if education:

                            for item in education:

                                if isinstance(
                                    item,
                                    dict,
                                ):

                                    degree = (
                                        item.get("degree")
                                        or ""
                                    )

                                    field = (
                                        item.get("field")
                                        or ""
                                    )

                                    institution = (
                                        item.get(
                                            "institution"
                                        )
                                        or ""
                                    )

                                    education_text = (
                                        " - ".join(
                                            value
                                            for value in [
                                                degree,
                                                field,
                                                institution,
                                            ]
                                            if value
                                        )
                                    )

                                    if education_text:

                                        st.write(
                                            f"• {education_text}"
                                        )

                                else:

                                    st.write(
                                        f"• {item}"
                                    )

                        else:

                            st.write(
                                "No education information "
                                "was extracted."
                            )

                        # =====================================
                        # PROJECTS
                        # =====================================

                        st.subheader("Projects")

                        projects = (
                            profile.get(
                                "projects",
                                [],
                            )
                        )

                        if projects:

                            for project in projects:

                                if isinstance(
                                    project,
                                    dict,
                                ):

                                    project_name = (
                                        project.get(
                                            "name"
                                        )
                                        or "Unnamed Project"
                                    )

                                    st.markdown(
                                        f"**{project_name}**"
                                    )

                                    description = (
                                        project.get(
                                            "description"
                                        )
                                    )

                                    if description:

                                        st.write(
                                            description
                                        )

                                    technologies = (
                                        project.get(
                                            "technologies",
                                            [],
                                        )
                                    )

                                    if technologies:

                                        st.caption(
                                            "Technologies: "
                                            + ", ".join(
                                                map(
                                                    str,
                                                    technologies,
                                                )
                                            )
                                        )

                                else:

                                    st.write(
                                        f"• {project}"
                                    )

                        else:

                            st.write(
                                "No projects were extracted."
                            )

                        # =====================================
                        # RAW EVIDENCE
                        # =====================================

                        st.divider()

                        # =====================================
                        # HUMAN DECISION
                        # =====================================

                        st.divider()

                        st.info(
                            "The information above is decision "
                            "support. The recruiter decides "
                            "whether the candidate proceeds "
                            "to interview."
                        )

                        # =====================================
                        # EXISTING INTERVIEW
                        # =====================================

                        try:

                            existing_interview = (
                                get_existing_interview_info(
                                    candidate_id,
                                    job_id,
                                    interview_round,
                                )
                            )

                        except Exception as error:

                            st.error(
                                "Could not check interview "
                                f"status: {error}"
                            )

                            existing_interview = None

                        st.divider()

                        st.subheader(
                            "Interview Status"
                        )

                        if existing_interview:

                            existing_status = (
                                existing_interview.get(
                                    "status"
                                )
                            )

                            if (
                                existing_status
                                == "submitted"
                            ):

                                st.success(
                                    "Completed"
                                )

                            else:

                                st.info(
                                    "Pending"
                                )

                            status_col1, status_col2 = (
                                st.columns(2)
                            )

                            status_col1.write(
                                "**Interviewer:**"
                            )

                            status_col1.write(
                                existing_interview.get(
                                    "interviewer_id"
                                )
                            )

                            status_col2.write(
                                "**Interview Round:**"
                            )

                            status_col2.write(
                                existing_interview.get(
                                    "interview_round"
                                )
                            )

                            st.caption(
                                "Interview Session ID: "
                                f"{existing_interview.get('id')}"
                            )

                            try:

                                InterviewSessionFactory = database_session()
                                interview_db = InterviewSessionFactory()

                                try:

                                    interview_details = get_interview_details(
                                        interview_db,
                                        existing_interview.get("id"),
                                    )

                                finally:

                                    interview_db.close()

                                if interview_details:

                                    questions = (
                                        interview_details.get("questions")
                                        or []
                                    )

                                    with st.expander(
                                        "View Interview Questions",
                                        expanded=False,
                                    ):

                                        st.caption(
                                            f"{len(questions)} interview questions"
                                        )

                                        for number, question in enumerate(
                                            questions,
                                            start=1,
                                        ):

                                            st.markdown(
                                                f"### Question {number}"
                                            )

                                            info_col1, info_col2 = (
                                                st.columns(2)
                                            )

                                            info_col1.write(
                                                "**Category:** "
                                                f"{question.get('category', 'N/A')}"
                                            )

                                            info_col2.write(
                                                "**Difficulty:** "
                                                f"{question.get('difficulty', 'N/A')}"
                                            )

                                            st.write(
                                                question.get("question")
                                                or "Question unavailable."
                                            )

                                            expected_concepts = (
                                                question.get(
                                                    "expected_concepts"
                                                )
                                                or []
                                            )

                                            if expected_concepts:

                                                st.markdown(
                                                    "**Expected Concepts**"
                                                )

                                                for concept in expected_concepts:

                                                    st.write(
                                                        f"• {concept}"
                                                    )

                                            criteria = (
                                                question.get(
                                                    "evaluation_criteria"
                                                )
                                                or []
                                            )

                                            if criteria:

                                                st.markdown(
                                                    "**Evaluation Criteria**"
                                                )

                                                for criterion in criteria:

                                                    st.write(
                                                        "• "
                                                        + str(
                                                            criterion.get(
                                                                "criterion",
                                                                "",
                                                            )
                                                        )
                                                    )

                                            st.divider()

                            except Exception as error:

                                st.error(
                                    "Could not load interview questions: "
                                    f"{error}"
                                )

                            

                        else:

                            st.warning(
                                "Not Started"
                            )

                        # =====================================
                        # COMPLETED EVALUATION
                        # =====================================

                        if (
                            existing_interview
                            and existing_interview.get(
                                "status"
                            )
                            == "submitted"
                        ):

                            try:

                                evaluation_report = (
                                    get_evaluation_report(
                                        existing_interview[
                                            "id"
                                        ],
                                        existing_interview[
                                            "interviewer_id"
                                        ],
                                    )
                                )

                                render_interview_evaluation(
                                    evaluation_report
                                )

                            except Exception as error:

                                st.error(
                                    "Could not load "
                                    "interview evaluation: "
                                    f"{error}"
                                )

                        # =====================================
                        # PROCEED TO INTERVIEW
                        # =====================================

                        if not existing_interview:

                            st.divider()

                            if active_interviewers:
                                if st.button("Assign interviewer", key=f"assign_{match_id}_{interview_round}", width="stretch"):
                                    @st.dialog(f"Assign interview · {candidate_name}", width="medium")
                                    def assign_interviewer_dialog():
                                        labels = {
                                            f"{user.full_name} (@{user.username}) · {user.email}": user
                                            for user in active_interviewers
                                        }
                                        selected = st.selectbox("Interviewer", list(labels), key=f"assign_user_{match_id}")
                                        round_choice = st.selectbox(
                                            "Interview round",
                                            ["Technical Round", "Technical + Project Round", "Final Technical Round"],
                                            key=f"assign_round_{match_id}",
                                        )
                                        if st.button("Confirm assignment", type="primary", key=f"confirm_assign_{match_id}", width="stretch"):
                                            chosen = labels[selected]
                                            st.session_state[assignment_key] = {
                                                "interviewer_id": chosen.id,
                                                "interviewer_email": chosen.email,
                                                "interviewer_name": chosen.full_name,
                                                "interview_round": round_choice,
                                            }
                                            st.rerun()
                                    assign_interviewer_dialog()

                            proceed_button = bool(assignment) or st.button(
                                "Proceed to Interview",
                                key=(
                                    f"proceed_{match_id}_"
                                    f"{interview_round}"
                                ),
                                type="primary",
                                use_container_width=True,
                            )

                            if proceed_button:

                                if not interviewer_id.strip():
                                    st.error(
                                        "No interviewer is available. "
                                        "Create an interviewer from User Management."
                                    )

                                else:

                                    try:

                                        duplicate = (
                                            get_existing_interview_info(
                                                candidate_id,
                                                job_id,
                                                interview_round,
                                            )
                                        )

                                        if duplicate:

                                            st.warning(
                                                "An interview "
                                                "already exists for "
                                                "this candidate, "
                                                "job and round."
                                            )

                                            st.rerun()

                                        # =========================
                                        # INTERVIEW AGENT
                                        # =========================

                                        with st.spinner(
                                            "Generating interview "
                                            "questions and "
                                            "evaluation criteria..."
                                        ):

                                            interview_report = (
                                                generate_interview_questions(
                                                    candidate_profile=(
                                                        profile
                                                    ),
                                                    job_profile=(
                                                        job_profile
                                                    ),
                                                    experience_report=(
                                                        experience_report
                                                    ),
                                                )
                                            )

                                            questions = (
                                                interview_report.get(
                                                    "questions",
                                                    [],
                                                )
                                            )

                                        if not questions:

                                            raise ValueError(
                                                "Interview Agent "
                                                "returned no "
                                                "questions."
                                            )

                                        if len(questions) != 6:

                                            raise ValueError(
                                                "Interview Agent must "
                                                "return exactly 6 "
                                                "questions. Received: "
                                                f"{len(questions)}"
                                            )

                                        # =========================
                                        # CREATE INTERVIEW
                                        # =========================

                                        InterviewSessionFactory = (
                                            database_session()
                                        )

                                        interview_db = (
                                            InterviewSessionFactory()
                                        )

                                        try:

                                            session_id = (
                                                create_interview(
                                                    interview_db,
                                                    candidate_id=(
                                                        candidate_id
                                                    ),
                                                    job_id=(
                                                        job_id
                                                    ),
                                                    interviewer_id=(
                                                        interviewer_id
                                                        .strip()
                                                    ),
                                                    target_job=(
                                                        job_title
                                                    ),
                                                    interview_round=(
                                                        interview_round
                                                    ),
                                                    candidate_summary={
                                                        "name": (
                                                            candidate_name
                                                        ),
                                                        "skills": (
                                                            profile.get(
                                                                "skills"
                                                            )
                                                            or []
                                                        ),
                                                        "education": (
                                                            profile.get(
                                                                "education"
                                                            )
                                                            or []
                                                        ),
                                                        "projects": (
                                                            profile.get(
                                                                "projects"
                                                            )
                                                            or []
                                                        ),
                                                        "experience": (
                                                            profile.get(
                                                                "experience"
                                                            )
                                                            or []
                                                        ),
                                                        "experience_level": (
                                                            experience_level
                                                        ),
                                                    },
                                                    questions=(
                                                        questions
                                                    ),
                                                )
                                            )

                                        finally:

                                            interview_db.close()

                                        st.session_state[
                                            "selected_candidate_id"
                                        ] = candidate_id

                                        st.session_state[
                                            "selected_job_id"
                                        ] = job_id

                                        st.session_state[
                                            "selected_match_id"
                                        ] = match_id

                                        st.session_state[
                                            "created_interview_"
                                            "session_id"
                                        ] = session_id

                                        st.success(
                                            "Interview created "
                                            f"successfully for "
                                            f"{candidate_name}."
                                        )

                                        st.write(
                                            "**Interview "
                                            "Session ID:**",
                                            session_id,
                                        )

                                        st.write(
                                            "**Interviewer ID:**",
                                            interviewer_id.strip(),
                                        )

                                        st.write(
                                            "**Questions "
                                            "Generated:**",
                                            len(questions),
                                        )

                                        try:
                                            send_interviewer_assignment_email(
                                                recipient=assignment.get("interviewer_email"),
                                                interviewer_name=assignment.get("interviewer_name", "Interviewer"),
                                                candidate_name=candidate_name,
                                                job_title=job_title,
                                                interview_round=interview_round,
                                            )
                                            st.session_state["assignment_email_notice"] = (
                                                "success",
                                                f"Assignment email sent to {assignment.get('interviewer_email')}.",
                                            )
                                        except Exception as mail_error:
                                            print(
                                                f"[email] FAILED to send assignment to "
                                                f"{assignment.get('interviewer_email') or 'the interviewer'}: "
                                                f"{mail_error}",
                                                flush=True,
                                            )
                                            st.session_state["assignment_email_notice"] = (
                                                "error",
                                                f"Interview was created, but email to {assignment.get('interviewer_email') or 'the interviewer'} failed: {mail_error}",
                                            )

                                        st.rerun()

                                    except Exception as error:

                                        st.error(
                                            "Unable to create "
                                            f"interview: {error}"
                                        )

                        # =====================================
                        # EXISTING INTERVIEW BUTTON
                        # =====================================

                        else:

                            st.divider()

                            if (
                                existing_interview.get(
                                    "status"
                                )
                                == "submitted"
                            ):

                                st.button(
                                    "Interview Completed",
                                    key=(
                                        f"completed_{match_id}_"
                                        f"{interview_round}"
                                    ),
                                    disabled=True,
                                    use_container_width=True,
                                )

                            else:

                                st.button(
                                    "Interview Already Created",
                                    key=(
                                        f"existing_{match_id}_"
                                        f"{interview_round}"
                                    ),
                                    disabled=True,
                                    use_container_width=True,
                                )
    finally:
        review_db.close()


with reports_tab:
    st.header("Interview performance", icon=":material/monitoring:")
    st.write(
        "Review completed interview outcomes across all candidates and job profiles. "
        "The workbook is rebuilt from saved evaluation reports whenever this page loads."
    )
    st.caption(
        "This report summarizes interviewer assessments; it does not make a hiring decision."
    )

    completed_interviews = []
    try:
        InterviewSessionFactory = database_session()
        interview_db = InterviewSessionFactory()
        try:
            records = interview_db.scalars(
                select(InterviewSession)
                .where(InterviewSession.status == "submitted")
                .order_by(InterviewSession.interview_date.desc())
            ).all()
            for record in records:
                report = dict(record.report.data or {}) if record.report else {}
                if not report:
                    continue
                summary = dict(record.candidate_summary or {})
                interview_date = record.interview_date
                completed_interviews.append({
                    "Interview ID": record.id,
                    "Candidate": summary.get("name") or "Unknown Candidate",
                    "Candidate ID": record.candidate_id,
                    "Job Profile": record.target_job or "Unspecified role",
                    "Job ID": record.job_id,
                    "Interview Round": record.interview_round,
                    "Interview Date": interview_date.isoformat() if interview_date else "",
                    "Interviewer ID": record.interviewer_id,
                    "Overall Score": report.get("overall_score", "Not assessed"),
                    "Technical Score": report.get("technical_score", "Not assessed"),
                    "Project Score": report.get("project_score", "Not assessed"),
                    "Behavioral Score": report.get("behavioral_score", "Not assessed"),
                    "Strengths": "; ".join(
                        str(item.get("criterion") or item.get("text") or item)
                        if isinstance(item, dict) else str(item)
                        for item in report.get("observed_strengths", [])
                    ),
                    "Areas to Improve": "; ".join(
                        str(item.get("criterion") or item.get("text") or item)
                        if isinstance(item, dict) else str(item)
                        for item in report.get("areas_requiring_improvement", [])
                    ),
                    "Report": report,
                })
        finally:
            interview_db.close()
    except Exception as error:
        st.error(f"Could not load completed interview reports: {error}")

    if not completed_interviews:
        st.info("No completed interview reports are available yet.")
    else:
        summary_rows = [
            {key: value for key, value in item.items() if key != "Report"}
            for item in completed_interviews
        ]
        detail_rows = []
        for item in completed_interviews:
            report = item["Report"]
            question_results = report.get("question_results") or report.get("questions") or []
            for question_number, result in enumerate(question_results, start=1):
                criteria = result.get("criteria") or []
                criteria = criteria or [{"criterion": "", "selected_option": ""}]
                for criterion in criteria:
                    detail_rows.append({
                        "Interview ID": item["Interview ID"],
                        "Candidate": item["Candidate"],
                        "Candidate ID": item["Candidate ID"],
                        "Job Profile": item["Job Profile"],
                        "Job ID": item["Job ID"],
                        "Interview Round": item["Interview Round"],
                        "Question": result.get("question_id", question_number),
                        "Category": result.get("category", ""),
                        "Question Text": result.get("question", ""),
                        "Checklist Criterion": criterion.get("criterion", ""),
                        "Assessment": criterion.get("selected_option", ""),
                        "Question Score": result.get("score", ""),
                    })

        job_profiles = sorted({row["Job Profile"] for row in summary_rows})
        job_filter_options = ["All job profiles", *job_profiles]
        filter_col, search_col = st.columns([1, 2])
        with filter_col:
            selected_profile = st.selectbox(
                "Filter preview by job profile",
                job_filter_options,
                key="performance_report_job_filter",
            )
        with search_col:
            candidate_search = st.text_input(
                "Search candidate or interviewer",
                placeholder="Enter a name or interviewer ID",
                key="performance_report_search",
            ).strip().casefold()

        preview_rows = [
            row for row in summary_rows
            if (selected_profile == "All job profiles" or row["Job Profile"] == selected_profile)
            and (
                not candidate_search
                or candidate_search in row["Candidate"].casefold()
                or candidate_search in row["Interviewer ID"].casefold()
            )
        ]
        st.caption(
            f"Showing {len(preview_rows)} of {len(summary_rows)} completed interview report(s)."
        )
        st.dataframe(
            pd.DataFrame(preview_rows),
            hide_index=True,
            width="stretch",
        )

        workbook_bytes = make_interview_performance_workbook(
            summary_rows,
            detail_rows,
        )
        st.download_button(
            "Download all job profiles as Excel",
            data=workbook_bytes,
            file_name="all_candidate_interview_performance.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
            icon=":material/download:",
            key="download_all_interview_performance",
        )

        with st.expander("View checklist-level performance"):
            if detail_rows:
                detail_ids = {row["Interview ID"] for row in preview_rows}
                visible_details = [
                    row for row in detail_rows
                    if row["Interview ID"] in detail_ids
                ]
                st.dataframe(
                    pd.DataFrame(visible_details),
                    hide_index=True,
                    width="stretch",
                )
            else:
                st.caption("No question-level checklist results were saved.")


# ============================================================
# USER MANAGEMENT
# ============================================================

with user_management_tab:

    st.header("User Management")

    st.write(
        "Create and manage interviewer accounts for the "
        "recruitment system."
    )

    st.info(
        "Only authenticated recruiters can create interviewer accounts."
    )

    # ========================================================
    # CREATE INTERVIEWER
    # ========================================================

    st.subheader("Create Interviewer")

    with st.form(
        "create_interviewer_form",
        clear_on_submit=True,
    ):

        col1, col2 = st.columns(2)

        with col1:

            interviewer_name = st.text_input(
                "Full Name",
                placeholder="Example: Interviewer Two",
            )

            interviewer_username = st.text_input(
                "Username",
                placeholder="Example: interviewer2",
            )

        with col2:

            interviewer_email = st.text_input(
                "Email",
                placeholder="Example: interviewer2@example.com",
            )

            interviewer_id_input = st.text_input(
                "Interviewer ID",
                placeholder="Example: INT002",
                help=(
                    "This ID is used when assigning interviews "
                    "to this interviewer."
                ),
            )

        password_col1, password_col2 = st.columns(2)

        with password_col1:

            interviewer_password = st.text_input(
                "Password",
                type="password",
            )

        with password_col2:

            interviewer_confirm_password = st.text_input(
                "Confirm Password",
                type="password",
            )

        create_interviewer_button = st.form_submit_button(
            "Create Interviewer",
            type="primary",
            use_container_width=True,
        )

    if create_interviewer_button:

        name = interviewer_name.strip()
        username = interviewer_username.strip()
        email = interviewer_email.strip()
        user_id = interviewer_id_input.strip().upper()

        if not name:

            st.error(
                "Full name is required."
            )

        elif not username:

            st.error(
                "Username is required."
            )

        elif not email:

            st.error(
                "Email is required."
            )

        elif not user_id:

            st.error(
                "Interviewer ID is required."
            )

        elif not user_id.startswith("INT"):

            st.error(
                "Interviewer ID should use the format INT001, INT002, etc."
            )

        elif not interviewer_password:

            st.error(
                "Password is required."
            )

        elif len(interviewer_password) < 8:

            st.error(
                "Password must contain at least 8 characters."
            )

        elif (
            interviewer_password
            != interviewer_confirm_password
        ):

            st.error(
                "Passwords do not match."
            )

        else:

            AuthSessionLocal = auth_database_session()

            auth_db = AuthSessionLocal()

            try:

                user = create_user(
                    auth_db,
                    user_id=user_id,
                    username=username,
                    email=email,
                    password=interviewer_password,
                    full_name=name,
                    role="interviewer",
                )

                st.success(
                    f"Interviewer {user.full_name} "
                    "created successfully."
                )

                st.write(
                    "**Interviewer ID:**",
                    user.id,
                )

                st.write(
                    "**Username:**",
                    user.username,
                )

                st.write(
                    "**Email:**",
                    user.email,
                )

            except ValueError as error:

                auth_db.rollback()

                st.error(
                    str(error)
                )

            except Exception as error:

                auth_db.rollback()

                st.error(
                    f"Could not create interviewer: {error}"
                )

            finally:

                auth_db.close()

    # ========================================================
    # EXISTING INTERVIEWERS
    # ========================================================

    st.divider()

    st.subheader("Registered Interviewers")

    try:

        AuthSessionLocal = auth_database_session()

        auth_db = AuthSessionLocal()

        try:

            interviewers = get_active_interviewers(
                auth_db
            )

        finally:

            auth_db.close()

        if not interviewers:

            st.info(
                "No active interviewers have been registered."
            )

        else:

            filter_col, sort_col = st.columns([2, 1])
            with filter_col:
                interviewer_search = st.text_input(
                    "Search registered interviewers",
                    placeholder="Name, username, email or ID",
                    key="registered_interviewer_search",
                ).strip().casefold()
            with sort_col:
                interviewer_sort = st.selectbox(
                    "Sort by",
                    ["Name", "Username", "Interviewer ID"],
                    key="registered_interviewer_sort",
                )

            filtered_interviewers = [
                interviewer
                for interviewer in interviewers
                if not interviewer_search
                or interviewer_search in str(interviewer.full_name or "").casefold()
                or interviewer_search in str(interviewer.username or "").casefold()
                or interviewer_search in str(interviewer.email or "").casefold()
                or interviewer_search in str(interviewer.id or "").casefold()
            ]
            sort_key = {
                "Name": lambda interviewer: str(interviewer.full_name or "").casefold(),
                "Username": lambda interviewer: str(interviewer.username or "").casefold(),
                "Interviewer ID": lambda interviewer: str(interviewer.id or "").casefold(),
            }[interviewer_sort]
            filtered_interviewers.sort(key=sort_key)

            st.caption(
                f"Showing {len(filtered_interviewers)} of {len(interviewers)} active interviewer(s)."
            )

            if not filtered_interviewers:
                st.info("No interviewers match that search.")

            for interviewer in filtered_interviewers:

                with st.container(border=True):

                    col1, col2, col3 = st.columns(
                        [2, 2, 1]
                    )

                    with col1:

                        st.markdown(
                            f"**{interviewer.full_name}**"
                        )

                        st.caption(
                            f"@{interviewer.username}"
                        )

                    with col2:

                        st.write(
                            interviewer.email
                        )

                    with col3:

                        st.write(
                            f"**{interviewer.id}**"
                        )

                    if st.button(
                        "Edit details",
                        key=f"edit_interviewer_{interviewer.id}",
                        icon=":material/edit:",
                    ):
                        st.session_state["editing_interviewer_id"] = interviewer.id
                        st.rerun()

                    if st.session_state.get("editing_interviewer_id") == interviewer.id:
                        with st.form(f"edit_interviewer_form_{interviewer.id}"):
                            edited_name = st.text_input(
                                "Full name", value=interviewer.full_name,
                                key=f"edit_name_{interviewer.id}",
                            )
                            edited_email = st.text_input(
                                "Email", value=interviewer.email,
                                key=f"edit_email_{interviewer.id}",
                            )
                            save_edit, cancel_edit = st.columns(2)
                            save_clicked = save_edit.form_submit_button(
                                "Save changes", type="primary", width="stretch"
                            )
                            cancel_clicked = cancel_edit.form_submit_button(
                                "Cancel", width="stretch"
                            )

                        if cancel_clicked:
                            st.session_state.pop("editing_interviewer_id", None)
                            st.rerun()

                        if save_clicked:
                            edit_factory = auth_database_session()
                            edit_db = edit_factory()
                            try:
                                update_interviewer_details(
                                    edit_db,
                                    interviewer.id,
                                    full_name=edited_name,
                                    email=edited_email,
                                )
                                st.session_state.pop("editing_interviewer_id", None)
                                st.success("Interviewer details updated.")
                                st.rerun()
                            except ValueError as error:
                                edit_db.rollback()
                                st.error(str(error))
                            except Exception as error:
                                edit_db.rollback()
                                st.error(f"Could not update interviewer details: {error}")
                            finally:
                                edit_db.close()

    except Exception as error:

        st.error(
            f"Could not load interviewers: {error}"
        )

