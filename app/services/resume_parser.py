import os
import re
from pypdf import PdfReader
from docx import Document


# ============================================================
# PDF EXTRACTION
# ============================================================

def extract_pdf_text(file_path: str) -> str:

    reader = PdfReader(file_path)

    text_parts = []

    for page_number, page in enumerate(reader.pages, start=1):

        try:
            text = page.extract_text()

            if text:
                text_parts.append(text)

        except Exception as error:
            print(
                f"Warning: Could not extract page "
                f"{page_number}: {error}"
            )

    return "\n".join(text_parts)


# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_docx_text(file_path: str) -> str:

    document = Document(file_path)

    text_parts = []

    # Normal paragraphs
    for paragraph in document.paragraphs:

        if paragraph.text.strip():
            text_parts.append(
                paragraph.text
            )

    # Resume information may also be stored inside tables
    for table in document.tables:

        for row in table.rows:

            row_values = []

            for cell in row.cells:

                value = cell.text.strip()

                if value:
                    row_values.append(value)

            if row_values:
                text_parts.append(
                    " | ".join(row_values)
                )

    return "\n".join(text_parts)


# ============================================================
# CLEAN RESUME TEXT
# ============================================================

def clean_resume_text(text: str) -> str:

    if not text:
        return ""

    # Normalize line endings
    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    # Replace tabs with spaces
    text = text.replace("\t", " ")

    # Remove repeated spaces
    text = re.sub(
        r"[ ]+",
        " ",
        text
    )

    # Remove spaces around line breaks
    text = re.sub(
        r" *\n *",
        "\n",
        text
    )

    # Avoid excessive blank lines
    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text
    )

    return text.strip()


# ============================================================
# MAIN RESUME PARSER
# ============================================================

def parse_resume(file_path: str) -> str:

    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"Resume not found: {file_path}"
        )

    extension = os.path.splitext(
        file_path
    )[1].lower()

    if extension == ".pdf":

        text = extract_pdf_text(
            file_path
        )

    elif extension == ".docx":

        text = extract_docx_text(
            file_path
        )

    else:

        raise ValueError(
            "Unsupported resume format. "
            "Only PDF and DOCX are supported."
        )

    cleaned_text = clean_resume_text(
        text
    )

    if not cleaned_text:

        raise ValueError(
            "No readable text could be extracted "
            "from the resume."
        )

    return cleaned_text